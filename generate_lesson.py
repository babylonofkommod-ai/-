#!/usr/bin/env python3
"""Premium teaching lesson. Python 3.10+.
Install: python -m pip install -r requirements-lesson.txt
Generate: python generate_lesson.py
Generate, commit and push: python generate_lesson.py --push
Georgia/Arial are fetched if not provided through --fonts-dir.
All source narration is kept verbatim in contiguous per-slide Notes.
"""
DATA = [{'title': 'Ваш Кристалл\nнезабываемости',
  'cards': ['Есть красивые женщины. Есть умные.', 'Есть сексуальные. Есть интересные.'],
  'quote': 'А есть незабываемые.',
  'notes': 'Вводный блок\n'
           '«Ваш Кристалл незабываемости»\n'
           'Почему одних женщин забывают, а других — нет?\n'
           '\n'
           'Есть красивые женщины. Есть умные. Есть сексуальные. Есть интересные.\n'
           'А есть незабываемые.'},
 {'title': 'Незабываемая женщина',
  'cards': ['незабываемая — это не обязательно самая красивая женщина в комнате.',
            'Не самая молодая. Не самая сексуальная.',
            'И даже не самая интересная в привычном понимании.'],
  'notes': 'И первое, что я хочу, чтобы вы поняли: незабываемая — это не обязательно самая красивая женщина '
           'в комнате.\n'
           'Не самая молодая. Не самая сексуальная.\n'
           'И даже не самая интересная в привычном понимании.\n'
           'Вы наверняка встречали женщин, которые вроде бы ничем невероятным не выделяются, но мужчины '
           'почему-то очень сильно на них западают.'},
 {'title': 'Почему?',
  'cards': ['сначала интерес', 'потом привыкание', 'потом охлаждение.'],
  'quote': '«хорошая женщина» и «незабываемая женщина» — не одно и то же.',
  'notes': 'И наоборот: потрясающая женщина — красивая, умная, успешная — а отношения с мужчинами почему-то '
           'постоянно заканчиваются одним и тем же: сначала интерес, потом привыкание, потом охлаждение.\n'
           'Почему?\n'
           'Потому что «хорошая женщина» и «незабываемая женщина» — не одно и то же.'},
 {'title': 'Мужчина запоминает\nне только вас',
  'cards': ['Он может забыть ресторан.', 'Он может даже забыть какие-то ваши слова.'],
  'quote': 'Он запоминает себя рядом с вами.',
  'notes': 'Мужчина запоминает не только вас\n'
           'Он запоминает себя рядом с вами.\n'
           'И это одна из главных мыслей всей программы.\n'
           'Через несколько лет мужчина может забыть, в каком платье вы пришли на свидание.\n'
           'Он может забыть ресторан. Он может даже забыть какие-то ваши слова.\n'
           'Но он гораздо дольше способен помнить: как ему было рядом с вами.'},
 {'title': 'Каким он себя чувствовал.',
  'cards': ['Было ли ему легко? Интересно?',
            'Хотелось ли ему становиться сильнее? Чувствовал ли он себя мужчиной?',
            'Смеялся ли? Удивлялся? Восхищался?'],
  'notes': 'Каким он себя чувствовал. Было ли ему легко? Интересно?\n'
           'Хотелось ли ему становиться сильнее? Чувствовал ли он себя мужчиной?\n'
           'Смеялся ли? Удивлялся? Восхищался?'},
 {'title': 'Научиться оставлять следы.',
  'cards': ['Наша задача гораздо интереснее.', 'Научиться оставлять следы.'],
  'quote': 'И поэтому наша задача — не научиться «правильно вести себя с мужчиной».',
  'notes': 'И поэтому наша задача — не научиться «правильно вести себя с мужчиной».\n'
           'Наша задача гораздо интереснее. Научиться оставлять следы.'},
 {'title': 'Как возникает\nнезабываемость',
  'cards': ['ГРАНЬ → СОПРИКОСНОВЕНИЕ', '→ СЛЕД → ПАМЯТЬ', '→ НЕЗАБЫВАЕМОСТЬ'],
  'accent_all': True,
  'notes': 'Как возникает незабываемость\n'
           'Сразу же покажу вам нашу главную формулу незабываемости.\n'
           'ГРАНЬ → СОПРИКОСНОВЕНИЕ → СЛЕД → ПАМЯТЬ → НЕЗАБЫВАЕМОСТЬ\n'
           'У каждой женщины есть разные грани.\n'
           'И когда какая-то ваша грань соприкасается с мужчиной, она оставляет след.'},
 {'title': 'Встреча закончилась.\nА лёгкость осталась.',
  'cards': ['Например, вы умеете создавать невероятную лёгкость.',
            'Вы встретились — посмеялись, поговорили, разошлись.'],
  'quote': 'Это след.',
  'notes': 'Например, вы умеете создавать невероятную лёгкость. Вы встретились — посмеялись, поговорили, '
           'разошлись.\n'
           'Встреча закончилась. А лёгкость осталась.\n'
           'Это след.'},
 {'title': 'Память о вас',
  'cards': ['Мужчина вспоминает вас — и вместе с вашим образом возвращается это ощущение.',
            'А когда таких следов становится много,'],
  'quote': 'появляется то, что я называю эффектом незабываемости.',
  'notes': 'Проходит несколько дней. Мужчина вспоминает вас — и вместе с вашим образом возвращается это '
           'ощущение.\n'
           'Так появляется память о вас. А когда таких следов становится много, появляется то, что я называю '
           'эффектом незабываемости.'},
 {'title': 'Одна грань — один след',
  'table': [['Лёгкость', 'след лёгкости'],
            ['Вдохновение', 'след вдохновения'],
            ['Интеллект', 'интеллектуальный след'],
            ['Индивидуальность', 'след уникальности']],
  'notes': 'Одна грань — один след\n'
           'Например\n'
           'Лёгкость → след лёгкости, Вдохновение → след вдохновения\n'
           'Интеллект → интеллектуальный след, Индивидуальность → след уникальности'},
 {'title': 'Почему одной красоты\nнедостаточно.',
  'table': [['Эстетика', 'эстетический след'],
            ['Достоинство', 'след уважения'],
            ['Контраст', 'след новизны']],
  'quote': 'Красота тоже способна оставить след. Но это один след.',
  'notes': 'Эстетика → эстетический след, Достоинство → след уважения\n'
           'Контраст → след новизны\n'
           'И вот теперь становится понятно, почему одной красоты недостаточно.\n'
           'Красота тоже способна оставить след. Но это один след.'},
 {'title': 'Совершенно другой\nобъём впечатления.',
  'cards': ['восхищение, интерес, лёгкость,',
            'уважение, любопытство, вдохновение, тепло,',
            'яркое воспоминание.'],
  'notes': 'А теперь представьте женщину, после которой мужчина одновременно уносит:\n'
           'восхищение, интерес, лёгкость, уважение, любопытство, вдохновение, тепло,\n'
           'яркое воспоминание.\n'
           'Это уже совершенно другой объём впечатления.'},
 {'title': 'Закон многогранности',
  'cards': ['Незабываемость редко создаётся одной суперспособностью.', 'Она создаётся комбинацией.'],
  'quote': 'Чем больше раскрытых граней — тем объёмнее образ женщины.',
  'notes': 'Закон многогранности\n'
           'Чем больше раскрытых граней — тем объёмнее образ женщины.\n'
           'И это второй важнейший принцип нашей программы.\n'
           'Незабываемость редко создаётся одной суперспособностью.\n'
           'Она создаётся комбинацией.'},
 {'title': 'Только одна грань',
  'cards': ['если женщина только красивая — мужчина достаточно быстро привыкает к красоте.',
            'Если она только лёгкая — прекрасно, но этого может быть мало.',
            'Только умная — тоже одна грань.'],
  'notes': 'Потому что если женщина только красивая — мужчина достаточно быстро привыкает к красоте.\n'
           'Если она только лёгкая — прекрасно, но этого может быть мало. Только умная — тоже одна грань.'},
 {'title': 'Она создаётся комбинацией.',
  'cards': ['она лёгкая, но глубокая.',
            'Нежная, но с характером.',
            'Женственная, но со своим миром.',
            'Спокойная, но иногда совершенно неожиданная.'],
  'notes': 'А теперь представьте: она лёгкая, но глубокая. Нежная, но с характером.\n'
           'Женственная, но со своим миром. Спокойная, но иногда совершенно неожиданная.'},
 {'title': 'Невозможно поместить\nв одну категорию',
  'cards': ['С ней можно смеяться',
            'и через двадцать минут разговаривать о вещах, о которых мужчина почти ни с кем не говорит.'],
  'quote': 'Вот такую женщину уже невозможно поместить в одну категорию.',
  'notes': 'С ней можно смеяться — и через двадцать минут разговаривать о вещах, о которых мужчина почти ни '
           'с кем не говорит.\n'
           'Вот такую женщину уже невозможно поместить в одну категорию.'}]

import argparse
from pathlib import Path
import subprocess
import hashlib
import urllib.request

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

NAME = 'Ваш Кристалл незабываемости'
BG = (26,26,26)
CARD = (40,40,40)
GOLD = (212,175,55)
WHITE = (255,255,255)
X, WIDTH = 4.5*72, 8.0*72
PAGE_W, PAGE_H = 13.333*72, 7.5*72
FONT_URLS = {
 'georgia.ttf': 'https://raw.githubusercontent.com/justrajdeep/fonts/master/Georgia.ttf',
 'georgiai.ttf': 'https://raw.githubusercontent.com/justrajdeep/fonts/master/Georgia%20Italic.ttf',
 'arial.ttf': 'https://raw.githubusercontent.com/justrajdeep/fonts/master/Arial.ttf',
}

def fonts(directory):
    directory.mkdir(parents=True,exist_ok=True)
    for filename,url in FONT_URLS.items():
        path=directory/filename
        if not path.is_file():
            with urllib.request.urlopen(url,timeout=45) as response:
                data=response.read()
            path.write_bytes(data)
    for face,filename in [('Georgia','georgia.ttf'),('GeorgiaItalic','georgiai.ttf'),('Arial','arial.ttf')]:
        pdfmetrics.registerFont(TTFont(face,str(directory/filename)))
        from fontTools.ttLib import TTFont as FontToolsFont
        cmap=FontToolsFont(directory/filename).getBestCmap()
        assert all(ord(char) in cmap for char in 'Кристалл незабываемости'), face

def wrap(text,face,size,width):
    lines=[]
    for paragraph in text.split('\n'):
        line=''
        for word in paragraph.split():
            test=(line+' '+word).strip()
            if pdfmetrics.stringWidth(test,face,size)>width and line:
                lines.append(line);line=word
            else:line=test
        lines.append(line)
    assert all(pdfmetrics.stringWidth(line,face,size)<=width for line in lines)
    return lines

def rgb(c):return RGBColor(*c)
def pdf_color(c):return tuple(channel/255 for channel in c)

class Renderer:
    def __init__(self,output):
        self.prs=Presentation()
        # Exact user-specified widescreen dimensions.
        self.prs.slide_width=Inches(13.333)
        self.prs.slide_height=Inches(7.5)
        self.prs.core_properties.title=NAME+' — обучающий урок'
        self.prs.core_properties.language='ru-RU'
        self.pdf=canvas.Canvas(str(output/(NAME+'.pdf')),pagesize=(PAGE_W,PAGE_H))
        self.pdf.setTitle(NAME+' — обучающий урок')
        self.output=output

    def rectangle(self,y,height,quote=False):
        sh=self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.5), Inches(y/72), Inches(8.0), Inches(height/72))
        sh.fill.solid();sh.fill.fore_color.rgb=rgb(BG if quote else CARD)
        if quote:
            sh.left=Inches(4.5)+Pt(.5);sh.width=Inches(8.0)-Pt(1)
            sh.line.color.rgb=rgb(GOLD);sh.line.width=Pt(.75)
        else:sh.line.fill.background()
        self.pdf.setFillColorRGB(*pdf_color(BG if quote else CARD))
        self.pdf.setStrokeColorRGB(*pdf_color(GOLD));self.pdf.setLineWidth(.75)
        self.pdf.rect(X+(.5 if quote else 0),PAGE_H-y-height,WIDTH-(1 if quote else 0),height,fill=1,stroke=int(quote))

    def text(self,lines,y,size=18,face='Arial',color=WHITE,padding=2):
        leading=size*1.24
        box=self.slide.shapes.add_textbox(Inches(4.5),Inches(y/72), Inches(8.0), Inches((len(lines)*leading+8)/72))
        tf=box.text_frame;tf.word_wrap=False
        tf.margin_left=tf.margin_right=Pt(padding);tf.margin_top=tf.margin_bottom=Pt(0)
        for i,value in enumerate(lines):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
            p.text=value;p.font.name='Georgia' if face.startswith('Georgia') else 'Arial'
            p.font.italic=face=='GeorgiaItalic';p.font.size=Pt(size);p.font.color.rgb=rgb(color)
            p.space_before=p.space_after=Pt(0);p.line_spacing=Pt(leading)
        self.pdf.setFont(face,size);self.pdf.setFillColorRGB(*pdf_color(color))
        asc=pdfmetrics.getAscent(face)*size/1000
        for i,value in enumerate(lines):self.pdf.drawString(X+padding,PAGE_H-y-asc-i*leading,value)

    def table_row(self,y,left,right):
        self.rectangle(y,57)
        box=self.slide.shapes.add_textbox(Inches(4.5),Inches((y+17)/72),Inches(8.0),Inches(31/72))
        tf=box.text_frame;tf.margin_left=tf.margin_right=Pt(14);tf.margin_top=tf.margin_bottom=Pt(0)
        p=tf.paragraphs[0];p.font.name='Arial';p.font.size=Pt(18)
        tabs=OxmlElement('a:tabLst');tab=OxmlElement('a:tab')
        tab.set('pos',str(Inches(2.95)));tab.set('algn','l');tabs.append(tab)
        p._p.get_or_add_pPr().append(tabs)
        for value,c in [(left+'\t',GOLD),('→ '+right,WHITE)]:
            run=p.add_run();run.text=value;run.font.name='Arial';run.font.size=Pt(18);run.font.color.rgb=rgb(c)
        for xx,value,c in [(X+14,left,GOLD),(X+14+2.95*72,'→ '+right,WHITE)]:
            assert xx+pdfmetrics.stringWidth(value,'Arial',18)<=X+WIDTH-14
            self.pdf.setFont('Arial',18);self.pdf.setFillColorRGB(*pdf_color(c))
            self.pdf.drawString(xx,PAGE_H-y-17-pdfmetrics.getAscent('Arial')*18/1000,value)

    def render(self,data,index,total):
        assert data['notes'].strip()
        assert len(data.get('cards',data.get('table',[])))<=4
        self.slide=self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.slide.background.fill.solid();self.slide.background.fill.fore_color.rgb=rgb(BG)
        self.pdf.setFillColorRGB(*pdf_color(BG));self.pdf.rect(0,0,PAGE_W,PAGE_H,fill=1,stroke=0)
        title=wrap(data['title'],'Georgia',34,WIDTH-4)
        assert len(title)<=2
        self.text(title,36,34,'Georgia',GOLD)
        y=max(145,36+len(title)*34*1.24+24)
        for card in data.get('cards',[]):
            lines=wrap(card,'Arial',18,WIDTH-28)
            height=len(lines)*18*1.24+24
            assert len(lines)<=3,'Split into another teaching slide instead of reducing font size.'
            self.rectangle(y,height)
            self.text(lines,y+11,18,color=GOLD if data.get('accent_all') else WHITE,padding=14)
            y+=height+10
        for left,right in data.get('table',[]):
            self.table_row(y,left,right);y+=67
        if data.get('quote'):
            y+=8
            lines=wrap(data['quote'],'GeorgiaItalic',18,WIDTH-28)
            height=len(lines)*18*1.24+24
            self.rectangle(y,height,quote=True)
            self.text(lines,y+11,18,'GeorgiaItalic',GOLD,14)
            y+=height
        assert y<=490,(index,y,'Split content; never shrink fonts.')
        self.text([f'{index:02d} / {total:02d}'],510,16)
        self.slide.notes_slide.notes_text_frame.text=data['notes']
        self.pdf.showPage()

    def save(self):
        self.prs.save(self.output/(NAME+'.pptx'));self.pdf.save()

def validate(output):
    import fitz
    prs=Presentation(output/(NAME+'.pptx'))
    assert prs.slide_width==Inches(13.333) and prs.slide_height==Inches(7.5)
    assert len(prs.slides)==len(DATA)
    lecture=' '.join(' '.join(item['notes'].split()) for item in DATA)
    assert hashlib.sha256(lecture.encode('utf-8')).hexdigest()=='618f688756018c7010acb9c11e7ecc09340c80c0537e3d2ecc10cba238f96ef4'
    for slide,data in zip(prs.slides,DATA):
        assert slide.notes_slide.notes_text_frame.text==data['notes']
        for shape in slide.shapes:
            assert shape.left>=Inches(4.5)
            if shape.has_text_frame and shape.text:
                assert shape.width==Inches(8.0)
                for p in shape.text_frame.paragraphs:
                    for run in p.runs:
                        assert (run.font.name or p.font.name) in ('Arial','Georgia')
                        size=run.font.size or p.font.size
                        assert size and size.pt in (16,18,34)
    document=fitz.open(output/(NAME+'.pdf'))
    assert len(document)==len(DATA)
    for page in document:
        pix=page.get_pixmap(matrix=fitz.Matrix(1,1),clip=fitz.Rect(0,0,X,PAGE_H),alpha=False)
        assert pix.samples==bytes(BG)*(pix.width*pix.height),'Speaker area must be perfectly blank.'
    print(f'Validated {len(DATA)} slides: exact dimensions, fonts, synchronized Notes, blank 4.5-inch speaker area.')

def push(output):
    names=[NAME+'.pptx',NAME+'.pdf','generate_lesson.py','requirements-lesson.txt']
    subprocess.run(['git','add','--',*names],cwd=output,check=True)
    changed=subprocess.run(['git','diff','--cached','--quiet'],cwd=output)
    if changed.returncode==1:
        subprocess.run(['git','commit','-m','Create premium 16-slide teaching lesson with synchronized narration'],cwd=output,check=True)
    elif changed.returncode!=0:raise RuntimeError('Cannot inspect staged files')
    subprocess.run(['git','push','origin','HEAD:main'],cwd=output,check=True)
    local=subprocess.check_output(['git','rev-parse','HEAD'],cwd=output,text=True).strip()
    remote=subprocess.check_output(['git','ls-remote','origin','refs/heads/main'],cwd=output,text=True).split()[0]
    assert local==remote
    print('Push verified:',local)

def main():
    parser=argparse.ArgumentParser(description='Generate a premium teaching lesson; optionally commit and push to main.')
    parser.add_argument('--output-dir',type=Path,default=Path(__file__).resolve().parent)
    parser.add_argument('--fonts-dir',type=Path,default=Path(__file__).resolve().parent/'work'/'fonts')
    parser.add_argument('--push',action='store_true',help='Commit output files and this generator, then push origin main.')
    args=parser.parse_args();args.output_dir.mkdir(parents=True,exist_ok=True)
    fonts(args.fonts_dir)
    renderer=Renderer(args.output_dir)
    for i,slide in enumerate(DATA,1):renderer.render(slide,i,len(DATA))
    renderer.save();validate(args.output_dir)
    if args.push:push(args.output_dir)

if __name__=='__main__':main()
